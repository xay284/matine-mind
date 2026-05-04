from flask import Flask, request, jsonify, send_file
from flask_cors import CORS
from rag_query import ask, call_ollama, OLLAMA_EXTRACTOR
import json
import os
import tempfile

app = Flask(__name__)
CORS(app)

# ──────────────────────────────────────────────
# ROUTE UPLOAD TDR (PDF / DOCX)
# ──────────────────────────────────────────────
@app.route("/upload_tdr", methods=["POST"])
def upload_tdr():
    if "file" not in request.files:
        return jsonify({"error": "Aucun fichier reçu"}), 400
    file = request.files["file"]
    ext  = os.path.splitext(file.filename)[1].lower()
    if ext not in (".pdf", ".docx", ".doc"):
        return jsonify({"error": "Format non supporté (PDF ou DOCX uniquement)"}), 400

    with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
        file.save(tmp.name)
        tmp_path = tmp.name

    try:
        from tdr_parser import process_tdr
        result = process_tdr(tmp_path, call_ollama, ask, OLLAMA_EXTRACTOR)
        return jsonify(result)
    except Exception as e:
        import traceback
        print("❌ ERREUR UPLOAD TDR :")
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500
    finally:
        if os.path.exists(tmp_path):
            os.unlink(tmp_path)


# ──────────────────────────────────────────────
# ROUTE PRINCIPALE : QUESTION → TABLEAU JSON
# ──────────────────────────────────────────────
@app.route("/ask", methods=["POST"])
def ask_route():
    data  = request.get_json()
    query = data.get("query", "")
    if not query:
        return jsonify({"error": "Aucune question fournie"}), 400

    try:
        result = ask(query, verbose=False)
        return jsonify({
            "answer" : result.get("answer", ""),
            "table"  : result.get("table", {"criteres": [], "projets": []})
        })
    except Exception as e:
        import traceback
        error_details = traceback.format_exc()
        print("❌ ERREUR COMPLÈTE :")
        print(error_details)          # ← affiche l'erreur exacte dans le terminal
        return jsonify({"error": str(e), "details": error_details}), 500



# ──────────────────────────────────────────────
# ROUTE EXPORT WORD (.docx)
# ──────────────────────────────────────────────
@app.route("/export/word", methods=["POST"])
def export_word():
    data  = request.get_json()
    table = data.get("table", {})
    query = data.get("query", "Recherche de projets")

    criteres = table.get("criteres", [])
    projets  = table.get("projets", [])

    if not projets:
        return jsonify({"error": "Aucun projet à exporter"}), 400

    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches, Cm
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_ALIGN_VERTICAL
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        BORDEAUX       = RGBColor(0x9C, 0x1F, 0x4A)
        BORDEAUX_LIGHT = RGBColor(0xF5, 0xE6, 0xEC)
        WHITE          = RGBColor(0xFF, 0xFF, 0xFF)
        GREEN          = RGBColor(0x2D, 0x8A, 0x4E)
        GRAY_TEXT      = RGBColor(0x55, 0x55, 0x55)

        def set_cell_bg(cell, hex_color: str):
            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd  = OxmlElement('w:shd')
            shd.set(qn('w:val'), 'clear')
            shd.set(qn('w:color'), 'auto')
            shd.set(qn('w:fill'), hex_color)
            tcPr.append(shd)

        def set_cell_border(cell, **kwargs):
            tc   = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right'):
                tag  = OxmlElement(f'w:{edge}')
                tag.set(qn('w:val'),   kwargs.get('val', 'single'))
                tag.set(qn('w:sz'),    kwargs.get('sz', '4'))
                tag.set(qn('w:space'), '0')
                tag.set(qn('w:color'), kwargs.get('color', 'CCCCCC'))
                tcBorders.append(tag)
            tcPr.append(tcBorders)

        doc     = Document()
        section = doc.sections[0]
        section.page_width    = Cm(29.7)
        section.page_height   = Cm(21.0)
        section.left_margin   = Cm(1.5)
        section.right_margin  = Cm(1.5)
        section.top_margin    = Cm(1.5)
        section.bottom_margin = Cm(1.5)

        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title_run = title_para.add_run("Matine Mind — Références Projets")
        title_run.bold           = True
        title_run.font.size      = Pt(16)
        title_run.font.color.rgb = BORDEAUX

        sub_para = doc.add_paragraph()
        sub_run  = sub_para.add_run(f"Recherche : {query}")
        sub_run.font.size      = Pt(10)
        sub_run.font.italic    = True
        sub_run.font.color.rgb = GRAY_TEXT
        doc.add_paragraph()

        fixed_cols = ["Projet", "Année", "Budget", "Client", "Pays"]
        all_cols   = fixed_cols + criteres
        n_cols     = len(all_cols)

        page_width_dxa = 15120
        projet_w = 3200
        annee_w  = 800
        budget_w = 1400
        client_w = 1800
        pays_w   = 1200
        remaining = page_width_dxa - projet_w - annee_w - budget_w - client_w - pays_w
        crit_w    = max(700, remaining // max(len(criteres), 1)) if criteres else 0
        col_widths = [projet_w, annee_w, budget_w, client_w, pays_w] + [crit_w] * len(criteres)

        table_doc = doc.add_table(rows=2 + len(projets), cols=n_cols)
        table_doc.style = 'Table Grid'

        # Ligne 0 : double header
        row0 = table_doc.rows[0]
        for i in range(4):
            set_cell_bg(row0.cells[i], "9C1F4A")

        if n_cols > 1:
            row0.cells[0].merge(row0.cells[3])
            p_proj = row0.cells[0].paragraphs[0]
            p_proj.clear()
            run = p_proj.add_run("Project")
            run.bold             = True
            run.font.color.rgb   = WHITE
            run.font.size        = Pt(10)
            p_proj.alignment     = WD_ALIGN_PARAGRAPH.LEFT
            set_cell_bg(row0.cells[0], "9C1F4A")

        if criteres:
            start_crit = 5
            end_crit   = n_cols - 1
            for i in range(start_crit, n_cols):
                set_cell_bg(row0.cells[i], "7F183B")
            if end_crit >= start_crit:
                row0.cells[start_crit].merge(row0.cells[end_crit])
                p_tor = row0.cells[start_crit].paragraphs[0]
                p_tor.clear()
                run = p_tor.add_run("Key subjects related to the TOR")
                run.bold           = True
                run.font.color.rgb = WHITE
                run.font.size      = Pt(10)
                p_tor.alignment    = WD_ALIGN_PARAGRAPH.CENTER
                set_cell_bg(row0.cells[start_crit], "7F183B")

        # Ligne 1 : en-têtes colonnes
        row1 = table_doc.rows[1]
        for j, col_name in enumerate(all_cols):
            cell = row1.cells[j]
            set_cell_bg(cell, "9C1F4A")
            p   = cell.paragraphs[0]
            p.clear()
            run = p.add_run(col_name)
            run.bold           = True
            run.font.color.rgb = WHITE
            run.font.size      = Pt(9)
            p.alignment        = WD_ALIGN_PARAGRAPH.CENTER
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Lignes projets
        for i, projet in enumerate(projets):
            row    = table_doc.rows[2 + i]
            bg_hex = "FFFFFF" if i % 2 == 0 else "FBF2F6"
            matching = projet.get("matching", {})

            values = [
                projet.get("titre",  "N/A"),
                projet.get("annee",  "N/A"),
                projet.get("budget", "N/A"),
                projet.get("client", "N/A"),
                projet.get("pays",   "N/A"),
            ]
            for j, val in enumerate(values):
                cell = row.cells[j]
                set_cell_bg(cell, bg_hex)
                p   = cell.paragraphs[0]
                p.clear()
                run = p.add_run(str(val))
                run.font.size = Pt(8)
                if j == 0:
                    run.bold           = True
                    run.font.color.rgb = BORDEAUX
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            for k, crit in enumerate(criteres):
                cell = row.cells[5 + k]
                set_cell_bg(cell, bg_hex)
                p    = cell.paragraphs[0]
                p.clear()
                score = matching.get(crit, 0)
                if isinstance(score, bool):
                    score = 100 if score else 0
                if score >= 40:
                    run = p.add_run("✓")
                    run.bold           = True
                    run.font.color.rgb = GREEN
                    run.font.size      = Pt(11)
                p.alignment             = WD_ALIGN_PARAGRAPH.CENTER
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

        # Largeurs colonnes
        for row in table_doc.rows:
            for j, cell in enumerate(row.cells):
                if j < len(col_widths):
                    tc   = cell._tc
                    tcPr = tc.get_or_add_tcPr()
                    tcW  = OxmlElement('w:tcW')
                    tcW.set(qn('w:w'),    str(col_widths[j]))
                    tcW.set(qn('w:type'), 'dxa')
                    tcPr.append(tcW)

        doc.add_paragraph()
        footer_para = doc.add_paragraph()
        footer_run  = footer_para.add_run(f"Généré par Matine Mind — {len(projets)} projet(s) trouvé(s)")
        footer_run.font.size      = Pt(9)
        footer_run.font.italic    = True
        footer_run.font.color.rgb = GRAY_TEXT

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        doc.save(tmp.name)
        tmp.close()

        return send_file(
            tmp.name,
            as_attachment=True,
            download_name="matine_mind_projets.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": f"Erreur export Word : {str(e)}"}), 500


# ──────────────────────────────────────────────
# ROUTE EXPORT PDF
# ──────────────────────────────────────────────
@app.route("/export/pdf", methods=["POST"])
def export_pdf():
    data  = request.get_json()
    table = data.get("table", {})
    query = data.get("query", "Recherche de projets")

    criteres = table.get("criteres", [])
    projets  = table.get("projets", [])

    if not projets:
        return jsonify({"error": "Aucun projet à exporter"}), 400

    try:
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib.enums import TA_CENTER, TA_LEFT

        BORDEAUX_RL = colors.HexColor("#9C1F4A")
        BORDEAUX_D  = colors.HexColor("#7F183B")
        BORDEAUX_L  = colors.HexColor("#F5E6EC")
        GREEN_RL    = colors.HexColor("#2D8A4E")
        GRAY_RL     = colors.HexColor("#555555")
        WHITE_RL    = colors.white
        STRIPE      = colors.HexColor("#FBF2F6")

        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pdf")

        doc = SimpleDocTemplate(
            tmp.name,
            pagesize    = landscape(A4),
            leftMargin  = 1.5 * cm,
            rightMargin = 1.5 * cm,
            topMargin   = 1.5 * cm,
            bottomMargin= 1.5 * cm,
        )

        styles      = getSampleStyleSheet()
        title_style = ParagraphStyle('title', fontSize=14, textColor=BORDEAUX_RL, fontName='Helvetica-Bold', spaceAfter=4)
        sub_style   = ParagraphStyle('sub',   fontSize=9,  textColor=GRAY_RL,    fontName='Helvetica-Oblique', spaceAfter=12)

        story = [
            Paragraph("Matine Mind — Références Projets", title_style),
            Paragraph(f"Recherche : {query}", sub_style),
        ]

        fixed_cols     = ["Projet", "Année", "Budget", "Client", "Pays"]
        all_cols       = fixed_cols + criteres
        n_cols         = len(all_cols)

        page_w   = landscape(A4)[0] - 3 * cm
        proj_w   = page_w * 0.24
        annee_w  = page_w * 0.06
        budget_w = page_w * 0.11
        cli_w    = page_w * 0.14
        pays_w   = page_w * 0.09
        rem_w    = page_w - proj_w - annee_w - budget_w - cli_w - pays_w
        crit_w   = rem_w / max(len(criteres), 1) if criteres else 0
        col_widths_pdf = [proj_w, annee_w, budget_w, cli_w, pays_w] + [crit_w] * len(criteres)

        cell_style = ParagraphStyle('cell',  fontSize=7, fontName='Helvetica',      leading=9)
        cell_bold  = ParagraphStyle('cellb', fontSize=7, fontName='Helvetica-Bold', textColor=BORDEAUX_RL, leading=9)
        hdr_style  = ParagraphStyle('hdr',   fontSize=8, fontName='Helvetica-Bold', textColor=WHITE_RL, alignment=TA_CENTER, leading=10)

        row0 = (
            [Paragraph("Project", hdr_style)] + [""] * 3 + [""] +
            ([Paragraph("Key subjects related to the TOR", hdr_style)] + [""] * (len(criteres) - 1) if criteres else [])
        )
        row1      = [Paragraph(c, hdr_style) for c in all_cols]
        data_rows = [row0, row1]

        for projet in projets:
            matching = projet.get("matching", {})
            row = [
                Paragraph(str(projet.get("titre",  "N/A")), cell_bold),
                Paragraph(str(projet.get("annee",  "N/A")), cell_style),
                Paragraph(str(projet.get("budget", "N/A")), cell_style),
                Paragraph(str(projet.get("client", "N/A")), cell_style),
                Paragraph(str(projet.get("pays",   "N/A")), cell_style),
            ]
            for crit in criteres:
                score = matching.get(crit, 0)
                if isinstance(score, bool):
                    score = 100 if score else 0
                row.append(Paragraph("✓" if score >= 40 else "", cell_style))
            data_rows.append(row)

        pdf_table = Table(data_rows, colWidths=col_widths_pdf, repeatRows=2)

        ts = TableStyle([
            ('BACKGROUND',     (0, 0), (4, 0),          BORDEAUX_RL),
            ('BACKGROUND',     (5, 0), (n_cols - 1, 0), BORDEAUX_D),
            ('TEXTCOLOR',      (0, 0), (-1, 0),          WHITE_RL),
            ('SPAN',           (0, 0), (4, 0)),
            ('BACKGROUND',     (0, 1), (-1, 1),          BORDEAUX_RL),
            ('TEXTCOLOR',      (0, 1), (-1, 1),          WHITE_RL),
            ('ALIGN',          (0, 0), (-1, 1),          'CENTER'),
            ('VALIGN',         (0, 0), (-1, 1),          'MIDDLE'),
            ('ALIGN',          (5, 2), (-1, -1),         'CENTER'),
            ('VALIGN',         (0, 0), (-1, -1),         'MIDDLE'),
            ('GRID',           (0, 0), (-1, -1),         0.3, colors.HexColor("#DDDDDD")),
            ('ROWBACKGROUNDS', (0, 2), (-1, -1),         [WHITE_RL, STRIPE]),
            ('TOPPADDING',     (0, 0), (-1, -1),         4),
            ('BOTTOMPADDING',  (0, 0), (-1, -1),         4),
            ('LEFTPADDING',    (0, 0), (-1, -1),         4),
            ('RIGHTPADDING',   (0, 0), (-1, -1),         4),
            ('TEXTCOLOR',      (5, 2), (-1, -1),         GREEN_RL),
            ('FONTNAME',       (5, 2), (-1, -1),         'Helvetica-Bold'),
            ('FONTSIZE',       (5, 2), (-1, -1),         10),
        ])

        if len(criteres) > 1:
            ts.add('SPAN', (5, 0), (n_cols - 1, 0))

        pdf_table.setStyle(ts)
        story.append(pdf_table)
        story.append(Spacer(1, 0.5 * cm))

        footer_style = ParagraphStyle('footer', fontSize=8, textColor=GRAY_RL, fontName='Helvetica-Oblique')
        story.append(Paragraph(f"Généré par Matine Mind — {len(projets)} projet(s) trouvé(s)", footer_style))

        doc.build(story)
        tmp.close()

        return send_file(
            tmp.name,
            as_attachment=True,
            download_name="matine_mind_projets.pdf",
            mimetype="application/pdf"
        )

    except Exception as e:
        import traceback
        traceback.print_exc()
        return jsonify({"error": f"Erreur export PDF : {str(e)}"}), 500


if __name__ == "__main__":
    app.run(debug=True, port=5000)